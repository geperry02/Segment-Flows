import tempfile
import unittest
from datetime import date
from pathlib import Path

import openpyxl

from segment_flows.config import SEGMENTS, loc_key
from segment_flows.report import Report, Row, parse_report
from segment_flows.transform import build
from segment_flows.workbook import write_workbook

FIXTURE = Path(__file__).parent / "fixtures" / "2026-09-01_Timely.xlsx"
SEP1 = date(2026, 9, 1)

# The 9/1/2026 Timely row from the team's template, in SEGMENTS column order.
EXPECTED_FLOWS = [
    -86797, -42533, 11031, -974230, -1612100, -1618111, -1424043, -1424043, -986085,
    -986968, -949703, -1046875, -1435284, -1667825, -1673700, -1673911, -1527887,
    -1515887, -1772241, -1293751, -1325883, -1415416, -1237371, -318090, -857117,
    -865228, -1588484, 245172, 245172,
]
EXPECTED_OPCAP = [
    -1063496, -1227996, 1656473, -1548296, -1773296, -1823296, -1793296, -1793296,
    -1153418, -1153418, -1552555, -1540955, -1853955, -1853955, -1828955, -1795000,
    -1766255, -1833147, -2294000, -1530000, -1390000, -1647968, -1663553, -1407553,
    -1645000, -905000, -2395165, 432882, 241058,
]
KEYS = [loc_key(seg, name) for _, seg, name in SEGMENTS]


class TemplateSample(unittest.TestCase):
    def test_sept1_timely_matches_template(self):
        report = parse_report(FIXTURE)
        self.assertEqual((report.gas_day.date(), report.cycle), (SEP1, "TIMELY"))
        flows, opcap, design, _ = build({(SEP1, "Timely"): report}, [SEP1])
        self.assertEqual([flows[(SEP1, "Timely")][k] for k in KEYS], EXPECTED_FLOWS)
        self.assertEqual([opcap[(SEP1, "Timely")][k] for k in KEYS], EXPECTED_OPCAP)
        self.assertEqual(design[SEP1][loc_key(527, "STA 527")], (707223, 2395165))

    def test_workbook_layout_and_rerun_overwrites(self):
        report = parse_report(FIXTURE)
        with tempfile.TemporaryDirectory() as tmp:
            out = Path(tmp) / "out.xlsx"
            for _ in range(2):  # second run must overwrite, not append
                write_workbook(out, [SEP1], *build({(SEP1, "Timely"): report}, [SEP1]))
            wb = openpyxl.load_workbook(out)
            self.assertEqual(wb.sheetnames, ["Flows", "Op Cap", "Design Cap", "Log"])
            ws = wb["Flows"]
            self.assertEqual([ws.cell(2, 4).value, ws.cell(3, 4).value, ws.cell(4, 4).value],
                             ["04", 394, "STA 319 EAST"])
            self.assertEqual(ws.cell(5, 2).value.date(), SEP1)
            self.assertEqual([ws.cell(5, 4 + i).value for i in range(len(KEYS))], EXPECTED_FLOWS)
            self.assertEqual(ws.max_row, 10)  # 6 cycle rows for one day
            self.assertEqual([wb["Design Cap"].cell(r, 3).value for r in (5, 6)], ["TD1", "TD2"])


def rep(day, locs):
    """locs: {key: (td1_sched or None, td2_sched or None)}; None = direction not posted."""
    out = {}
    for key, (s1, s2) in locs.items():
        out[key] = {}
        if s1 is not None:
            out[key]["TD1"] = Row(design=1000, operating=900, scheduled=s1)
        if s2 is not None:
            out[key]["TD2"] = Row(design=2000, operating=800, scheduled=s2)
    return Report(day, "", "", out)


class DirectionRules(unittest.TestCase):
    K = KEYS[0]

    def run_days(self, days):
        """days: {date: {cycle label: (td1, td2)}} for KEYS[0]."""
        reports = {(d, c): rep(d, {self.K: v}) for d, cycles in days.items() for c, v in cycles.items()}
        return build(reports, sorted(days))

    def test_zero_flow_uses_day_majority(self):
        d = SEP1
        flows, opcap, _, _ = self.run_days({d: {"Timely": (0, 0), "Eve": (0, 50), "ID1": (0, 60),
                                                "ID2": (40, 0)}})
        self.assertEqual(opcap[(d, "Timely")][self.K], -800)
        self.assertEqual(flows[(d, "Timely")][self.K], 0)
        self.assertEqual(opcap[(d, "ID2")][self.K], 900)

    def test_all_zero_day_uses_8am_history(self):
        d1, d2, d3 = date(2026, 9, 1), date(2026, 9, 2), date(2026, 9, 3)
        _, opcap, _, _ = self.run_days({
            d1: {"8AM": (0, 10)}, d2: {"8AM": (0, 10)}, d3: {"Timely": (0, 0), "8AM": (0, 0)},
        })
        self.assertEqual(opcap[(d3, "Timely")][self.K], -800)

    def test_no_history_defaults_to_td1_and_logs(self):
        _, opcap, _, notes = self.run_days({SEP1: {"Timely": (0, 0)}})
        self.assertEqual(opcap[(SEP1, "Timely")][self.K], 900)
        self.assertTrue(any("defaulted to TD1" in n[4] for n in notes))

    def test_single_direction_keeps_its_sign(self):
        _, opcap, _, _ = self.run_days({SEP1: {"Timely": (None, 0), "Eve": (None, 5)}})
        self.assertEqual(opcap[(SEP1, "Timely")][self.K], -800)

    def test_flow_both_directions_is_logged(self):
        flows, _, _, notes = self.run_days({SEP1: {"Timely": (10, 30)}})
        self.assertEqual(flows[(SEP1, "Timely")][self.K], -30)
        self.assertTrue(any("both directions" in n[4] for n in notes))


if __name__ == "__main__":
    unittest.main()
