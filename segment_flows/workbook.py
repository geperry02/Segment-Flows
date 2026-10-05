"""Write the cleaned data into the Flows / Op Cap / Design Cap workbook layout.

Layout (matches the team template):
  row 1 blank, row 2 Zone, row 3 Segment, row 4 Date / Loc header + location names,
  data from row 5: column B date (first row of each day), column C cycle,
  column D onward one column per segment location.
"""

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font

from .config import CYCLE_LABELS, SEGMENTS, loc_key

FIRST_COL = 4
FIRST_ROW = 5
NUM_FMT = "#,##0"
DATE_FMT = "m/d/yyyy"
BOLD = Font(bold=True)


def _header(ws):
    ws["C2"], ws["C3"], ws["B4"], ws["C4"] = "Zone", "Segment ", "Date", "Loc"
    for i, (zone, seg, name) in enumerate(SEGMENTS):
        col = FIRST_COL + i
        ws.cell(2, col, zone)
        ws.cell(3, col, seg)
        ws.cell(4, col, name).alignment = Alignment(wrap_text=True, horizontal="center")
        ws.column_dimensions[ws.cell(1, col).column_letter].width = 13
    for row in (2, 3, 4):
        for cell in ws[row]:
            cell.font = BOLD
    ws.column_dimensions["B"].width = 11
    ws.row_dimensions[4].height = 45
    ws.freeze_panes = ws.cell(FIRST_ROW, FIRST_COL)


def _sheet(wb, title):
    if title in wb.sheetnames:
        return wb[title]
    ws = wb.create_sheet(title)
    _header(ws)
    return ws


def _row_index(ws, labels):
    """{(date, label): row} for rows already in the sheet, so re-runs overwrite in place."""
    index, current = {}, None
    for r in range(FIRST_ROW, ws.max_row + 1):
        date, label = ws.cell(r, 2).value, ws.cell(r, 3).value
        if date is not None:
            current = date.date() if hasattr(date, "date") else date
        if current is not None and label in labels:
            index[(current, label)] = r
    return index


def _write_block(ws, dates, labels, values):
    """values: {(date, label): {key: number}}; adds rows for missing dates."""
    keys = [loc_key(seg, name) for _, seg, name in SEGMENTS]
    index = _row_index(ws, labels)
    next_row = max(index.values(), default=FIRST_ROW - 1) + 1
    for day in dates:
        for i, label in enumerate(labels):
            r = index.get((day, label))
            if r is None:
                r, next_row = next_row, next_row + 1
                if i == 0:
                    ws.cell(r, 2, day).number_format = DATE_FMT
                ws.cell(r, 3, label)
            for j, key in enumerate(keys):
                v = values.get((day, label), {}).get(key)
                cell = ws.cell(r, FIRST_COL + j, v)
                cell.number_format = NUM_FMT


def write_workbook(path, dates, flows, opcap, design, notes):
    try:
        wb = load_workbook(path)
    except FileNotFoundError:
        wb = Workbook()
        wb.remove(wb.active)

    _write_block(_sheet(wb, "Flows"), dates, CYCLE_LABELS, flows)
    _write_block(_sheet(wb, "Op Cap"), dates, CYCLE_LABELS, opcap)

    design_values = {}
    for day, per_key in design.items():
        for key, (td1, td2) in per_key.items():
            design_values.setdefault((day, "TD1"), {})[key] = td1
            design_values.setdefault((day, "TD2"), {})[key] = -td2 if td2 is not None else None
    _write_block(_sheet(wb, "Design Cap"), dates, ["TD1", "TD2"], design_values)

    if "Log" in wb.sheetnames:
        del wb["Log"]
    log = wb.create_sheet("Log")
    log.append(["Date", "Cycle", "Segment", "Location", "Note"])
    for cell in log[1]:
        cell.font = BOLD
    for day, cycle, seg, name, msg in notes:
        log.append([day, cycle, seg, name, msg])
        log.cell(log.max_row, 1).number_format = DATE_FMT
    for col, width in zip("ABCDE", (11, 8, 9, 26, 90)):
        log.column_dimensions[col].width = width

    wb.save(path)
